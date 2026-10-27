"""
pipeline/fundae.py — FUNDAE (Fundación Estatal para la Formación en el Empleo): series estadísticas de formación.

FUNDAE no publica una API ni ficheros de datos: sus «Series estadísticas» (www.fundae.es › Datos estadísticos) se ven en
un visor de SQL Server Reporting Services (visor?reportPath=rptSeriesNN). Este lector abre el visor con su propia sesión,
toma la URL de exportación (ExportUrlBase) y descarga cada serie en Excel.

El servidor falla a menudo: devuelve una serie distinta de la pedida (el visor guarda el informe en el servidor y se cruzan
las peticiones), o un 404/500 al exportar. Por eso:
  - cada Excel se acepta solo si el nombre de su hoja es el de la serie pedida (rptSeriesNN) y su título es el esperado;
  - cada serie se guarda por separado (datos/normalizado/fundae/rptSeriesNN.csv, en la caché entre ejecuciones) y la tabla
    se monta con la última versión buena de cada una: las que fallan hoy se completan en las cargas siguientes;
  - el tiempo total está limitado para no retrasar el resto de la carga.

Tabla (formato largo): serie; cuadro; anyo; dimension; categoria; subdimension; subcategoria; ccaa_id; medida; valor; url
  medida: Participantes · Horas de formación · Horas de permiso · Empresas formadoras · Crédito dispuesto
  ccaa_id: código de comunidad (C01…C19) cuando la categoría es una comunidad autónoma
"""
import html as html_mod, http.cookiejar, io, re, time, unicodedata, urllib.parse, urllib.request

BASE = "https://www.fundae.es/"
UA = {"User-Agent": "Mozilla/5.0 rrhh360-pipeline/1.0"}

# serie -> (palabras que tiene que tener el título, dimensión de las categorías, subdimensión de las columnas o None)
SERIES = {
    "01": ("por edad. Total", "Edad", None),
    "07": ("extranjeros", "Edad (extranjeros)", None),
    "02": ("grupo de cotización", "Grupo de cotización", None),
    "03": ("tamaño de la empresa", "Tamaño de la empresa", None),
    "04": ("nivel de estudios. Total", "Nivel de estudios", None),
    "09": ("colectivos prioritarios", "Colectivo prioritario", None),
    "10": ("modalidad de impartición", "Modalidad", None),
    "14": ("familia profesional", "Familia profesional", None),
    "15": ("área profesional", "Área profesional", None),
    "13": ("por contenidos formativos", "Contenidos formativos", None),  # la 11 trae lo mismo y el visor la sirve peor
    "30": ("Comunidad Autónoma del participante", "Comunidad autónoma", None),
    "20": ("beneficiarios por edad", "Edad (permisos de formación)", None),
    "22": ("beneficiarios por nivel de estudios", "Nivel de estudios (permisos de formación)", None),
    "23": ("objeto del PIF", "Titulación de acceso (permisos de formación)", None),
    "48": ("beneficiarios por Comunidad Autónoma", "Comunidad autónoma (permisos de formación)", "Sexo"),
    "38": ("según tamaño de la empresa", "Tamaño de la empresa (empresas formadoras)", None),
    "41": ("organización de la formación", "Organización de la formación", "Tamaño de la empresa"),
    "49": ("Comunidad Autónoma y tamaño", "Comunidad autónoma de la empresa", "Tamaño de la empresa"),
    "50": ("Comunidad Autónoma y sector", "Comunidad autónoma de la empresa", "Sector de actividad"),
}
MEDIDAS = {"participantes": "Participantes", "participantes formados": "Participantes", "horas de formacion": "Horas de formación",
           "horas de permiso": "Horas de permiso", "empresas formadoras": "Empresas formadoras",
           "credito dispuesto": "Crédito dispuesto"}
CCAA = {"andalucia": "C01", "aragon": "C02", "asturiasprincipadode": "C03", "balearsilles": "C04", "canarias": "C05",
        "cantabria": "C06", "castillayleon": "C07", "castillalamancha": "C08", "cataluna": "C09",
        "comunitatvalenciana": "C10", "extremadura": "C11", "galicia": "C12", "madridcomunidadde": "C13",
        "murciaregionde": "C14", "navarracomunidadforalde": "C15", "paisvasco": "C16", "riojala": "C17", "ceuta": "C18",
        "melilla": "C19"}
CABECERA = ["serie", "cuadro", "anyo", "dimension", "categoria", "subdimension", "subcategoria", "ccaa_id", "medida",
            "valor", "url"]


def _clave(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", s)


def _url(url):
    p = urllib.parse.urlsplit(url)
    return urllib.parse.urlunsplit((p.scheme, p.netloc, urllib.parse.quote(urllib.parse.unquote(p.path)),
                                    urllib.parse.quote(urllib.parse.unquote(p.query), safe="=&[],"), ""))


def exportar(serie, contexto, intentos=3, pausa=5):
    """Excel (bytes) de la serie, comprobado (hoja rptSeriesNN y título esperado). Lanza ValueError si no se consigue."""
    import openpyxl
    visor = f"{BASE}publicaciones/visor?reportPath=rptSeries{serie}"
    ultimo = ""
    for i in range(intentos):
        if i:
            time.sleep(pausa * i)
        op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()),
                                         urllib.request.HTTPSHandler(context=contexto))
        op.addheaders = list(UA.items())
        try:
            with op.open(_url(visor), timeout=90) as r:
                html = r.read().decode("utf-8", "replace")
            src = re.findall(r'<iframe[^>]+src="([^"]+)"', html, re.I)
            if not src:
                ultimo = "el visor no tiene el informe"; continue
            rv_url = urllib.parse.urljoin(visor, html_mod.unescape(src[0]))
            with op.open(_url(rv_url), timeout=120) as r:
                rv, rv_final = r.read().decode("utf-8", "replace"), r.geturl()
            time.sleep(2)  # el informe se termina de generar en el servidor tras cargar el visor
            exp = re.findall(r'"ExportUrlBase"\s*:\s*"([^"]+)"', rv)
            if not exp:
                ultimo = "sin URL de exportación"; continue
            base = urllib.parse.urljoin(rv_final, exp[0].replace("\\u0026", "&"))
            if f"FileName=rptSeries{serie}&" not in base:
                ultimo = "el visor ha devuelto otra serie (" + (re.findall(r"FileName=(\w+)", base) or ["?"])[0] + ")"; continue
            with op.open(base + "EXCELOPENXML", timeout=180) as r:
                datos = r.read()
            if datos[:2] != b"PK":
                ultimo = "la exportación no es un Excel"; continue
            ws = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True).worksheets[0]
            titulo = next((str(v) for f in ws.iter_rows(max_row=6, values_only=True) for v in f if v), "")
            if ws.title != f"rptSeries{serie}" or _clave(SERIES[serie][0]) not in _clave(titulo):
                ultimo = f"el Excel es de otra serie ({ws.title}: {titulo[:60]})"; continue
            return datos
        except Exception as e:
            ultimo = f"{type(e).__name__}: {e}"[:200]
    raise ValueError(f"FUNDAE rptSeries{serie}: {ultimo}")


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", ".")) if re.fullmatch(r"-?[\d.,]+", str(v).strip()) else None
    except ValueError:
        return None


def _anyo(v):
    s = str(v or "").strip()
    return int(s) if re.fullmatch(r"(19|20)\d\d", s) else None


def leer(serie, datos, url=""):
    """Excel de una serie -> filas en formato largo (CABECERA)."""
    import openpyxl
    _, dimension, subdim = SERIES[serie]
    ws = openpyxl.load_workbook(io.BytesIO(datos), data_only=True).worksheets[0]
    filas_x = [list(f) for f in ws.iter_rows(values_only=True)]
    cuadro = next((str(v).strip() for f in filas_x for v in f if v), "")
    anyo_fijo = next((int(m.group(1)) for f in filas_x for v in f if v
                      for m in [re.match(r"A[ñn]o:\s*(\d{4})", str(v))] if m), None)
    # cabecera: primera fila con 2 o más textos a partir de la columna 3 (categorías o años)
    cab, i_cab = {}, None
    for i, f in enumerate(filas_x):
        textos = {j: str(v).strip() for j, v in enumerate(f) if j >= 3 and isinstance(v, str) and v.strip()}
        if len(textos) >= 2:
            cab, i_cab = textos, i
            break
    if not cab:
        raise ValueError(f"rptSeries{serie}: no se encuentra la cabecera")
    columnas_anyo = all(_anyo(t) for t in cab.values())
    out, bloque, anyo_bloque = [], None, None
    primera = min(cab)  # las etiquetas (año, categoría, medida) van en las columnas anteriores a los datos
    for f in filas_x[i_cab + 1:]:
        etiquetas = [str(v).strip() for v in f[:primera] if isinstance(v, (str, int)) and str(v).strip()]
        if any(e.startswith("NOTAS") for e in etiquetas):
            break
        valores = {j: _num(f[j]) for j in cab if j < len(f) and f[j] not in (None, "")}
        valores = {j: v for j, v in valores.items() if v is not None}
        medida, etiqueta = None, None
        for e in etiquetas:
            if _clave(e) in {_clave(k) for k in MEDIDAS}:
                medida = MEDIDAS[next(k for k in MEDIDAS if _clave(k) == _clave(e))]
            else:
                etiqueta = e
        if etiqueta is not None and _anyo(etiqueta) and not columnas_anyo:
            anyo_bloque, etiqueta = _anyo(etiqueta), None
        if not valores:
            if etiqueta is not None and medida is None and etiqueta not in ("Año",) and not etiqueta.startswith("CC.AA") \
                    and not etiqueta.startswith("Comunidad") and not etiqueta.startswith("Familia"):
                bloque = etiqueta
            continue
        if medida is None:  # filas sin medida (rptSeries41): la etiqueta es la categoría y el valor, empresas
            medida, categoria = "Empresas formadoras", etiqueta
        else:
            if etiqueta is not None:
                bloque = etiqueta
            categoria = bloque
        for j, v in valores.items():
            col = cab[j]
            if columnas_anyo:
                anyo, cat, sub = _anyo(col), categoria, ""
            elif subdim:
                anyo, cat, sub = anyo_fijo or anyo_bloque, categoria, col
            else:
                anyo, cat, sub = anyo_fijo or anyo_bloque, col, ""
            if anyo is None or cat is None:
                continue
            out.append([serie, cuadro, anyo, dimension, cat, subdim or "", sub, CCAA.get(_clave(cat), ""), medida, v, url])
    if not out:
        raise ValueError(f"rptSeries{serie}: sin datos")
    return out


def cargar(contexto, carpeta, limite_s=420, pausa=1.0):
    """Intenta actualizar cada serie (hasta `limite_s` segundos en total) y monta la tabla con la última versión buena de
    cada una, guardada en `carpeta`. Devuelve (CABECERA, filas, estado {serie: 'hoy' | 'anterior' | 'sin datos'}).
    Lanza ValueError si no hay ninguna serie (ni de hoy ni de cargas anteriores)."""
    import csv
    carpeta.mkdir(parents=True, exist_ok=True)
    t0, estado, errores = time.time(), {}, {}
    for serie in SERIES:
        if time.time() - t0 > limite_s:
            errores[serie] = "sin tiempo en esta carga"
            continue
        url = f"{BASE}publicaciones/visor?reportPath=rptSeries{serie}"
        try:
            filas_s = leer(serie, exportar(serie, contexto), url)
            with open(carpeta / f"rptSeries{serie}.csv", "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh); w.writerow(CABECERA); w.writerows(filas_s)
            estado[serie] = "hoy"
        except Exception as e:
            errores[serie] = f"{type(e).__name__}: {e}"[:200]
        time.sleep(pausa)
    filas = []
    for serie in SERIES:
        p = carpeta / f"rptSeries{serie}.csv"
        if p.exists():
            with open(p, encoding="utf-8") as fh:
                filas += [f for f in csv.reader(fh)][1:]
            estado.setdefault(serie, "anterior")
        else:
            estado[serie] = "sin datos"
    if not filas:
        raise ValueError("FUNDAE: ninguna serie disponible · " + " · ".join(f"{k}: {v}" for k, v in errores.items())[:600])
    return CABECERA, filas, estado
