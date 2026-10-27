#!/usr/bin/env python3
"""
pipeline/fuentes.py — Fase 2: lectores de SEPE, Seguridad Social y Ministerio de Trabajo.

Cada lector descarga la fuente, la lleva a una tabla limpia y la deja en datos/normalizado/<tabla>.csv:
  sepe_paro_municipio       SEPE · paro registrado por municipio y mes (sexo x edad, sector)        [CSV anual]
  sepe_contratos_municipio  SEPE · contratos registrados por municipio y mes (sexo x tipo, sector)  [CSV anual]
  sepe_demandantes_municipio SEPE · demandantes de empleo por municipio y mes (sexo x edad, sector)  [CSV anual]
  sepe_demandantes_nivel    SEPE · demandantes parados por nivel formativo, sexo y nacionalidad (nacional, un fichero por mes)
  ss_it_detalle             Seguridad Social · IT por mes, régimen, contingencia, sexo y provincia (duración media)
  ss_afiliacion             Seguridad Social · afiliación (PX-Web): provincia/CCAA, régimen, CNAE, contrato, sexo
  ss_incapacidad_temporal   Seguridad Social · IT por provincia: duración media, incidencia y prevalencia por mil protegidos
  mites_cpd_actividad       Ministerio de Trabajo · contratos de puesta a disposición por mes y sección CNAE (ETT-E2)
  mites_convenios           Ministerio de Trabajo · convenios: variación salarial pactada, jornada (CCT-2.1)
  mites_accidentes          Ministerio de Trabajo · accidentes en jornada e índices de incidencia por CCAA y provincia,
                            acumulado del año del último avance (ATR-I1.1)
  mites_ett_mensual         ETT · CPD por modalidad, cesiones, trabajadores por sexo/nacionalidad/edad, contratos (E1/E3/E4)
  mites_ett_territorio      ETT · cesiones y CPD por modalidad por CCAA y provincia, mes y acumulado (ETT-1)
  mites_ett_actividad       ETT · CPD por modalidad por sector y división CNAE (ETT-2)
  mites_ett_resumen         ETT · acumulado por sección, CCAA, perfil y afiliados, con variación oficial interanual (R1..R4)
  mites_convenios_detalle   Convenios · por sector, CCAA/provincia, ámbito, tramo de jornada y cláusula de garantía
  mites_accidentes_detalle  Accidentes · gravedad, provincia, actividad, ocupación, sexo y edad, índices y variación (ATR)
  mites_regulacion_empleo   ERE/ERTE · trabajadores afectados por tipo, sexo, actividad, territorio, causa y tamaño (10 meses)
  fundae_formacion          FUNDAE · series estadísticas de formación para el empleo: participantes, horas, empresas
                            formadoras y crédito por edad, estudios, grupo de cotización, tamaño, CCAA, familia… (anual)
  eurostat_indicadores      Eurostat · población 16-29, tasa de empleo 15-29, NEET por nivel de formación y por urbanización
  (los lectores de detalle están en pipeline/mites_detalle.py y el de FUNDAE en pipeline/fundae.py)
Calidad: datos/qa/fuentes_carga.csv (estado, filas y periodos de cada lector).
Si un lector falla y existe la última versión buena de su tabla (conservada en caché entre ejecuciones), se mantiene
y se marca como ANTERIOR, con la fecha de esa carga (datos/estado_cargas.json).

Uso:  python pipeline/fuentes.py [--solo sepe ss mites] [--desde 2019]
Requiere: pip install openpyxl xlrd cryptography
Los valores ocultos por secreto estadístico ('<5' en el SEPE) quedan vacíos y se cuentan en n_secretos.
"""
import argparse, hashlib, csv, io, re, sys, time, unicodedata, urllib.parse, urllib.request
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
from catalogo_fuentes import contexto  # cadena TLS de mites.gob.es sin desactivar la verificación
import estado
import eurostat as eurostat_lector
import fundae as fundae_lector
import mites_detalle

NORM, QA = RAIZ / "datos" / "normalizado", RAIZ / "datos" / "qa"
CFG = RAIZ / "pipeline" / "config"
UA = {"User-Agent": "Mozilla/5.0 rrhh360-pipeline/1.0"}
MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
         "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
MESES_ABR = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "oct": 10,
             "nov": 11, "dic": 12}

SEPE = "https://sede.sepe.gob.es/es/portaltrabaja/resources/sede/datos_abiertos/datos/{}_por_municipios_{}_csv.csv"
SEPE_NF = "https://sede.sepe.gob.es/es/portaltrabaja/resources/sede/datos_abiertos/datos/Demandantes_empleo_por_NF_{}_csv.csv"
SS_CONSULTAS = {  # consultas guardadas de PX-Web (vistas en datos.gob.es, publicador E05072501)
    "afil_medios_prov_cnae09_regimen": "https://w6.seg-social.es/PXWeb/sq/e7afe15e-d110-4584-907c-b25235dd3f42",
    "afil_ultimo_dia_prov_cnae": "https://w6.seg-social.es/PXWeb/sq/d921596c-f52d-478b-94b8-0998230648a6",
    "afil_medios_rg_prov_contrato_sexo": "https://w6.seg-social.es/PXWeb/sq/adef39b2-bd54-47f5-a387-e08d372c0825",
    "afil_medios_ccaa_cnae25_regimen": "https://w6.seg-social.es/PXWeb/sq/b97adb6b-47a9-439b-ae2d-f6693ff5cd12",
}
MITES_PAGINAS = {  # página de cada estadística: de aquí se toma el fichero más reciente
    "ett": ("https://www.mites.gob.es/es/estadisticas/mercado_trabajo/ETT/welcome.htm", r"ETT_(\d{1,2})_(\d{4})\.xlsx$"),
    "cct": ("https://www.mites.gob.es/es/estadisticas/condiciones_trabajo_relac_laborales/CCT/welcome.htm",
            r"CCT_(\d{1,2})_(\d{4})\.xlsx?$"),
    "atr": ("https://www.mites.gob.es/es/estadisticas/condiciones_trabajo_relac_laborales/EAT/welcome.htm",
            r"ATR_(\d{1,2})_(\d{4})\.xlsx$"),
    "reg": ("https://www.mites.gob.es/es/estadisticas/condiciones_trabajo_relac_laborales/REG/welcome.htm",
            r"reg_(\d{1,2})_(\d{4})\.xlsx?$"),
}


# ------------------------------------------------------------------------------------------------ utilidades
_CAIDAS = {}  # host -> descargas seguidas que han fallado del todo (tras todos los intentos)


def descargar(url, intentos=3, timeout=60):
    """Hasta 3 intentos de 60 s (tiempo máximo sin recibir datos, no de la descarga entera): la web del Ministerio de Trabajo corta a veces las conexiones desde GitHub y
    esperar 5 minutos por intento dejaba el paso bloqueado (13 ficheros de ERTE, cierres de convenios…).
    Si un servidor falla del todo dos veces seguidas, el resto de descargas a ese servidor se cortan al momento:
    cada fuente afectada usa su última carga buena en vez de esperar otros 5 minutos por fichero."""
    host = urllib.parse.urlparse(url).hostname
    if _CAIDAS.get(host, 0) >= 2:
        raise TimeoutError(f"{host} no responde en esta ejecución: no se reintenta")
    for i in range(intentos):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout,
                                        context=contexto(host)) as r:
                datos = r.read()
            _CAIDAS[host] = 0
            return datos
        except urllib.error.HTTPError as e:
            if e.code == 404:
                raise
            if i == intentos - 1:
                _CAIDAS[host] = _CAIDAS.get(host, 0) + 1
                raise
        except Exception:
            if i == intentos - 1:
                _CAIDAS[host] = _CAIDAS.get(host, 0) + 1
                raise
        time.sleep(10 * (i + 1))


def texto(b):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue


def num(v):
    """'1234', '1234.5', 1234.0 -> float; '<5', '', None -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s.startswith("<") or s in ("-", ".", ".."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def norm_nombre(s):
    """Para casar nombres de territorio escritos de forma distinta: sin tildes, sin signos, palabras ordenadas."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = s.replace("c. foral", "comunidad foral").replace("s. c. ", "santa cruz de ")
    palabras = [p for p in re.split(r"[^a-z]+", s) if p and p not in ("de", "la", "el", "las", "los", "y")]
    return " ".join(sorted(palabras))


def cargar_territorios():
    with open(CFG / "territorio.csv", encoding="utf-8") as f:
        filas = list(csv.DictReader(f, delimiter=";"))
    por_nombre = {}
    for t in filas:
        for n in t["nombre"].split("/") + [t["nombre"]]:
            por_nombre.setdefault((t["nivel"], norm_nombre(n)), t["territorio_id"])
    return por_nombre


def escribir(nombre, cabecera, filas):
    NORM.mkdir(parents=True, exist_ok=True)
    with open(NORM / f"{nombre}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(cabecera); w.writerows(filas)


_FICHEROS = {}


def ficheros_mites(clave):
    """{(año, mes): url} de todos los avances mensuales enlazados en la página de la estadística."""
    if clave not in _FICHEROS:
        pagina, patron = MITES_PAGINAS[clave]
        html = texto(descargar(pagina))
        encontrados = {}
        for href in re.findall(r'href="([^"]+)"', html):
            m = re.search(patron, href, re.I)
            if m:
                encontrados[(int(m.group(2)), int(m.group(1)))] = urllib.parse.urljoin(pagina, href)
        if not encontrados:
            raise ValueError(f"no se encontró ningún fichero {patron} en {pagina}")
        _FICHEROS[clave] = encontrados
    return _FICHEROS[clave]


def ultimo_fichero(clave):
    k, url = max(ficheros_mites(clave).items())
    return url, k


_DESCARGAS = {}


CACHE_MITES = RAIZ / "datos" / "raw" / "mites"  # se conserva entre ejecuciones (caché del workflow)


def descargar_mites(url, fijo=False):
    """Descarga una vez por ejecución (varios lectores leen hojas distintas del mismo fichero).
    fijo=True: fichero de un periodo cerrado (cierre de diciembre de convenios, meses anteriores de ERTE). Cada
    periodo tiene su propia URL, así que se guarda en disco y no se vuelve a descargar en las cargas siguientes."""
    if url not in _DESCARGAS:
        disco = CACHE_MITES / hashlib.sha1(url.encode()).hexdigest()
        if fijo and disco.exists():
            _DESCARGAS[url] = disco.read_bytes()
        else:
            _DESCARGAS[url] = descargar(url)
            if fijo:
                CACHE_MITES.mkdir(parents=True, exist_ok=True)
                disco.write_bytes(_DESCARGAS[url])
    return _DESCARGAS[url]


def hoja(datos, nombre):
    """Devuelve las filas (listas de valores) de una hoja de Excel .xlsx o .xls."""
    if datos[:4] == b"\xd0\xcf\x11\xe0":
        import xlrd
        h = xlrd.open_workbook(file_contents=datos).sheet_by_name(nombre)
        return [h.row_values(i) for i in range(h.nrows)]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    nombre_real = next(n for n in wb.sheetnames if n.strip() == nombre.strip())
    return [list(r) for r in wb[nombre_real].iter_rows(values_only=True)]


def celda(fila, i):
    v = fila[i] if i < len(fila) else None
    return "" if v is None else str(v).strip()


# ------------------------------------------------------------------------------------------------ SEPE
def sepe(desde, solo=None):
    resultado = []
    for tipo, tabla in (("Paro", "sepe_paro_municipio"), ("Contratos", "sepe_contratos_municipio"),
                        ("Dtes_empleo", "sepe_demandantes_municipio")):
        if solo and tabla != solo:
            continue
        filas, cab_valores, anyos = [], None, []
        for anyo in range(desde, date.today().year + 1):
            try:
                t = texto(descargar(SEPE.format(tipo, anyo)))
            except urllib.error.HTTPError as e:
                print(f"  SEPE {tipo} {anyo}: HTTP {e.code} (se omite)"); continue
            lineas = t.splitlines()
            i0 = next(i for i, l in enumerate(lineas) if l.lower().startswith("código mes") or l.lower().startswith("codigo mes"))
            cab = [c.strip() for c in lineas[i0].split(";")]
            valores = [slug(c) for c in cab[8:] if c]
            if cab_valores is None:
                cab_valores = valores
            elif valores != cab_valores:
                raise ValueError(f"SEPE {tipo} {anyo}: columnas distintas a las de años anteriores: {valores}")
            for l in lineas[i0 + 1:]:
                c = l.split(";")
                if len(c) < 9 or not re.fullmatch(r"\d{6}", c[0].strip()):
                    continue
                v = [c[j] if j < len(c) else "" for j in range(8, 8 + len(cab_valores))]
                n_secretos = sum(1 for x in v if x.strip().startswith("<"))
                cm = c[0].strip()
                filas.append([f"{cm[:4]}-{cm[4:]}-01", "M" + c[6].strip().zfill(5), c[7].strip(),
                              "P" + c[4].strip().zfill(2), "C" + c[2].strip().zfill(2), n_secretos] + [num(x) for x in v])
            anyos.append(anyo)
        escribir(tabla, ["fecha", "territorio_id", "municipio", "provincia_id", "ccaa_id", "n_secretos"] + cab_valores, filas)
        fechas = [f[0] for f in filas]
        resultado.append([tabla, "OK", len(filas), min(fechas, default=""), max(fechas, default=""),
                          f"años {anyos}; columnas: {', '.join(cab_valores)}"])
    return resultado


# ------------------------------------------------------------------------------------------------ Seguridad Social
# la sección es el primer "X. TEXTO EN MAYÚSCULAS" de la cabecera (el régimen puede empezar por letra y punto: "R. General")
RX_SECCION = re.compile(r"^(.*?)\s*((?:\b[A-V]\.\s[A-ZÁÉÍÓÚÑ\"]{2}.*)|TODAS LAS SECCIONES)$")  # el nombre de la sección puede llevar puntos ("E. … SANEAMIENTO. GESTIÓN …")
RX_SEXO = re.compile(r"^(.*?)\s+(Total Sexo|Mujer|Mujeres|Varón|Varones|Hombres)$", re.I)


def parsear_nivel_formativo(t, fecha):
    """Texto del CSV del SEPE «Demandantes de empleo por nivel formativo» de un mes -> filas
    [fecha, nivel, hombres, mujeres, espanoles, extranjeros, total]. Las cifras vienen con punto de millares («392.725»).
    Las columnas de antigüedad no se leen: sus cifras no cuadran con otras estadísticas (casi todo «más de 12 meses»)."""
    def n(v):
        v = str(v).strip().strip('"').replace(".", "")
        return int(v) if re.fullmatch(r"\d+", v) else None
    lineas = [l for l in t.splitlines() if l.strip()]
    if not lineas or "nivel formativo" not in lineas[0].lower():
        raise ValueError("cabecera inesperada (no es el fichero por nivel formativo)")
    cab = [c.strip().strip('"').lower() for c in lineas[0].split(";")]
    idx = {k: next((i for i, c in enumerate(cab) if k in c), None) for k in ("hombres", "mujeres", "españoles", "extranjeros")}
    if None in idx.values():
        raise ValueError(f"faltan columnas esperadas: {cab}")
    filas = []
    for l in lineas[1:]:
        c = [x.strip().strip('"') for x in l.split(";")]
        h, m = n(c[idx["hombres"]]), n(c[idx["mujeres"]])
        if not c[0] or h is None or m is None:
            continue
        filas.append([fecha, c[0], h, m, n(c[idx["españoles"]]), n(c[idx["extranjeros"]]), h + m])
    if len(filas) < 5:
        raise ValueError(f"solo {len(filas)} niveles")
    return filas


def sepe_nivel_formativo(desde=2019):
    """SEPE: un fichero nacional por mes (Demandantes_empleo_por_NF_AAAAMM). Los meses sin fichero se omiten."""
    cab = ["fecha", "nivel", "hombres", "mujeres", "espanoles", "extranjeros", "total"]
    filas, meses = [], []
    hoy = date.today()
    for a in range(desde, hoy.year + 1):
        for m in range(1, 13):
            if (a, m) > (hoy.year, hoy.month):
                break
            try:
                t = texto(descargar(SEPE_NF.format(f"{a}{m:02d}")))
                filas += parsear_nivel_formativo(t, f"{a}-{m:02d}-01")
                meses.append(f"{a}-{m:02d}")
            except urllib.error.HTTPError:
                continue
            except ValueError as e:
                print(f"  SEPE nivel formativo {a}-{m:02d}: se omite ({e})")
    if not filas:
        raise ValueError("SEPE nivel formativo: ningún mes descargado")
    escribir("sepe_demandantes_nivel", cab, filas)
    return [["sepe_demandantes_nivel", "OK", len(filas), meses[0], meses[-1], f"{len(meses)} meses"]]


def ss():
    filas, resumen = [], []
    for consulta, url in SS_CONSULTAS.items():
        lineas = texto(descargar(url)).splitlines()
        i0 = next(i for i, l in enumerate(lineas) if l.startswith('"Periodo"'))
        rows = list(csv.reader(lineas[i0:], delimiter=";"))
        cab, nivel = rows[0], ("Provincia" if "provincia" in rows[0][1].lower() else "CCAA")
        columnas = []
        for h in cab[2:]:
            sexo = modalidad = seccion = regimen = None
            resto = h.strip()
            m = RX_SEXO.match(resto)
            if m:  # consulta por modalidad de contrato y sexo ("INDEFINIDO-T. Completo Mujer")
                modalidad, sexo = m.group(1), m.group(2)
            elif RX_SECCION.match(resto):  # consulta por régimen y sección CNAE
                regimen, seccion = RX_SECCION.match(resto).group(1) or None, RX_SECCION.match(resto).group(2)
            else:
                regimen = resto
            columnas.append((regimen, seccion, modalidad, sexo))
        n = 0
        for r in rows[1:]:
            if len(r) < 3 or not re.fullmatch(r"\d{6}", r[0].strip()):
                continue
            per, terr = r[0].strip(), r[1].strip()
            m = re.match(r"^(\d+)\.\s", terr)
            tid = ("P" if nivel == "Provincia" else "C") + m.group(1).zfill(2) if m else (
                "N00" if terr.lower().startswith("total nacional") else None)
            for (regimen, seccion, modalidad, sexo), v in zip(columnas, r[2:]):
                filas.append([f"{per[:4]}-{per[4:]}-01", tid, terr if tid is None else "", consulta,
                              regimen, seccion, modalidad, sexo, num(v)])
                n += 1
        fechas = sorted({r[0].strip() for r in rows[1:] if r and re.fullmatch(r"\d{6}", r[0].strip())})
        resumen.append([f"ss_afiliacion:{consulta}", "OK", n, fechas[0] if fechas else "", fechas[-1] if fechas else "",
                        f"{len(columnas)} columnas, nivel {nivel}"])
    escribir("ss_afiliacion", ["fecha", "territorio_id", "territorio_sin_resolver", "consulta", "regimen", "seccion_cnae",
                               "modalidad_contrato", "sexo", "afiliados"], filas)
    return resumen


SS_IT_INDICE = "https://www.seg-social.es/wps/portal/wss/internet/EstadisticasPresupuestosEstudios/Estadisticas/EST45/EST46"


_IT_DETALLE, _IT_DETALLE_CAB = [], []  # lo rellena ss_it() al leer los mismos ficheros


def ss_it_detalle():
    """Detalle de la IT (mismos ficheros que ss_it): se escribe con lo que ss_it ya ha leído."""
    if not _IT_DETALLE:
        ss_it()
    if not _IT_DETALLE:
        raise ValueError("sin detalle de IT")
    escribir("ss_it_detalle", _IT_DETALLE_CAB, _IT_DETALLE)
    periodos = sorted({f"{f[0]}-{f[1]:02d}" for f in _IT_DETALLE})
    return [["ss_it_detalle", "OK", len(_IT_DETALLE), periodos[0], periodos[-1], "duración, incidencia y prevalencia por mes, "
             "régimen, contingencia, sexo y provincia"]]


def ss_it():
    """Incapacidad temporal por CCAA y provincia (Seguridad Social): fichero del sistema del último ejercicio
    (enero al último mes publicado) y cierre del ejercicio anterior. Se localizan en las páginas 'Ejercicio AAAA'."""
    import html as html_mod
    indice = texto(descargar(SS_IT_INDICE))
    ejercicios = {}
    for href, txt in re.findall(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', indice, re.S):
        m = re.search(r"Ejercicio\s+(\d{4})", re.sub(r"<[^>]+>", "", txt))
        if m:
            ejercicios[int(m.group(1))] = urllib.parse.urljoin(SS_IT_INDICE, html_mod.unescape(href))
    if not ejercicios:
        raise ValueError("no se encuentran las páginas 'Ejercicio AAAA' de incapacidad temporal")
    filas, cab, usados = [], None, []
    for anyo in sorted(ejercicios)[-2:]:
        pagina = texto(descargar(ejercicios[anyo]))
        enlaces = [html_mod.unescape(h) for h in re.findall(r'href="([^"]+\.xlsx[^"]*)"', pagina)]
        sistema = [h for h in enlaces if re.search(r"sistema", urllib.parse.unquote(h), re.I)]
        cierre = [h for h in sistema if re.search(r"cierre", urllib.parse.unquote(h), re.I)]
        elegido = (cierre or sistema or [None])[0]
        if not elegido:
            print(f"  IT {anyo}: sin fichero del sistema en la página"); continue
        url = urllib.parse.urljoin("https://www.seg-social.es/", elegido)
        datos = descargar(url)
        cab, f = mites_detalle.it_provincias(datos, url.split("?")[0])
        filas += f
        try:  # detalle por mes, régimen, contingencia, sexo y provincia (si falla, no impide la tabla principal)
            cab_d, fd = mites_detalle.it_detalle(datos, url.split("?")[0])
            _IT_DETALLE.extend(fd)
            _IT_DETALLE_CAB[:] = cab_d
        except Exception as e:
            print(f"  IT detalle {anyo}: {type(e).__name__}: {e}")
        usados.append(f"{anyo}: {urllib.parse.unquote(url.split('/')[-1].split('?')[0])}")
    if not filas:
        raise ValueError("ningún fichero de IT legible")
    escribir("ss_incapacidad_temporal", cab, filas)
    periodos = sorted({f"{f[0]}-{f[1]:02d}" for f in filas})
    return [["ss_incapacidad_temporal", "OK", len(filas), periodos[0], periodos[-1], " · ".join(usados)]]


# ------------------------------------------------------------------------------------------------ Ministerio de Trabajo
def mites_ett():
    url, (anyo_f, mes_f) = ultimo_fichero("ett")
    filas_hoja = hoja(descargar_mites(url), "ETT-E2")
    # La hoja reparte las secciones CNAE en varios bloques verticales, cada uno con su propia cabecera
    # ("AÑOS" + fila de letras): en cada cabecera se rehace el mapa columna -> sección.
    filas, anyo, cols, vistas = [], None, None, set()
    for i, f in enumerate(filas_hoja):
        c0 = celda(f, 0)
        if c0.upper().startswith("AÑOS"):
            letras = filas_hoja[i + 1] if i + 1 < len(filas_hoja) else []
            cols, anyo = [], None
            for j in range(1, max(len(f), len(letras))):
                if celda(f, j).upper().startswith("TOTAL"):
                    cols.append((j, "TOTAL", "Total"))
                elif re.fullmatch(r"[A-V](\s*y\s*[A-V])?", celda(letras, j)):  # "A" o agrupadas: "P y V", "T y U"
                    cols.append((j, re.sub(r"\s+", " ", celda(letras, j)), celda(f, j)))
            continue
        if cols is None:
            continue
        m = re.match(r"^(\d{4})\s*(\(\*\))?\s*:$", c0)
        if m:
            anyo = int(m.group(1)); continue
        mes = MESES.get(c0.lower())
        if anyo and mes:
            for j, cod, nom in cols:
                v = num(f[j] if j < len(f) else None)
                clave = (anyo, mes, cod)
                if v is not None and clave not in vistas:
                    vistas.add(clave)
                    filas.append([f"{anyo}-{mes:02d}-01", cod, nom, v, url])
    escribir("mites_cpd_actividad", ["fecha", "seccion_cnae", "seccion_nombre", "cpd", "fichero"], filas)
    fechas = [f[0] for f in filas]
    secciones = sorted({f[1] for f in filas} - {"TOTAL"})
    # control: la suma de secciones frente al TOTAL de cada mes (el total incluye "los no clasificados", nota (1))
    tot = {f[0]: f[3] for f in filas if f[1] == "TOTAL"}
    suma = {}
    for f in filas:
        if f[1] != "TOTAL":
            suma[f[0]] = suma.get(f[0], 0) + f[3]
    desv = {k: 100 * (tot[k] - suma.get(k, 0)) / tot[k] for k in tot if tot[k]}
    max_desv = max((abs(v) for v in desv.values()), default=0)
    estado = "OK" if max_desv <= 0.5 else "AVISO"
    return [["mites_cpd_actividad", estado, len(filas), min(fechas, default=""), max(fechas, default=""),
             f"{url} · {len(secciones)} secciones ({', '.join(secciones)}) · máxima diferencia entre el total y la suma "
             f"de secciones: {max_desv:.2f}% (no clasificados)"]]


def mites_cct():
    url, _ = ultimo_fichero("cct")
    filas_hoja = hoja(descargar_mites(url), "CCT-2.1")
    filas, bloque_anyo, en_meses = [], None, False
    for f in filas_hoja:
        c0 = celda(f, 0)
        m = re.match(r"^(\d{4})(?:\.0)?\s*(\(.*)?$", c0)
        if m and num(f[1] if len(f) > 1 else None) is not None:  # fila anual con datos
            filas.append([int(m.group(1)), None, "Año completo" if not m.group(2) else "Año (provisional)",
                          num(f[1]), num(f[2]), num(f[3]), num(f[4]), num(f[5]), num(f[6]), url])
            continue
        if m:  # cabecera de bloque mensual ("2025 (4) (*)")
            bloque_anyo, en_meses = int(m.group(1)), False; continue
        if c0.lower().startswith("registrados hasta"):
            en_meses = True; continue
        mes = MESES_ABR.get(c0.lower().rstrip(".")[:3])
        if en_meses and bloque_anyo and mes:
            filas.append([bloque_anyo, mes, "Registrados hasta el mes", num(f[1]), num(f[2]), num(f[3]), num(f[4]),
                          num(f[5]), num(f[6]), url])
    escribir("mites_convenios", ["anyo_efectos", "mes_registro", "tipo_dato", "convenios", "empresas", "trabajadores",
                                 "variacion_salarial_pactada_pct", "variacion_salarial_revisada_pct",
                                 "jornada_media_horas_anyo", "fichero"], filas)
    anyos = [f[0] for f in filas]
    return [["mites_convenios", "OK", len(filas), min(anyos, default=""), max(anyos, default=""), url]]


def mites_atr():
    url, (anyo, mes) = ultimo_fichero("atr")
    territorios = cargar_territorios()
    filas, sin_casar = [], []
    for f in hoja(descargar_mites(url), "ATR-I1.1"):
        nombre = celda(f, 0)
        total = num(f[2] if len(f) > 2 else None)
        if not nombre or total is None:
            continue
        if nombre.upper() == "TOTAL":
            tid = "N00"
        else:
            nivel = "CCAA" if nombre.isupper() else "Provincia"
            tid = territorios.get((nivel, norm_nombre(nombre)))
            if tid is None:  # comunidades uniprovinciales que aparecen en mayúsculas: se intenta como provincia
                tid = territorios.get(("Provincia", norm_nombre(nombre))) if nivel == "CCAA" else None
            if tid is None:
                sin_casar.append(nombre)
        filas.append([f"{anyo}-01-01", f"{anyo}-{mes:02d}-01", tid, nombre, total, num(f[3]), num(f[5]), num(f[7]),
                      num(f[8]), url])
    escribir("mites_accidentes", ["periodo_desde", "periodo_hasta", "territorio_id", "territorio_nombre",
                                  "accidentes_jornada", "accidentes_mortales", "afiliados_media_mensual",
                                  "indice_incidencia", "indice_incidencia_mortales", "fichero"], filas)
    return [["mites_accidentes", "OK" if not sin_casar else "AVISO", len(filas), f"{anyo}-01", f"{anyo}-{mes:02d}",
             f"{url}" + (f" · territorios sin casar: {sin_casar}" if sin_casar else "")]]


def mites_ett_detalle(tabla):
    """ETT del último avance: serie mensual por modalidad y perfil (ETT-E1/E3/E4), detalle por territorio (ETT-1) y por
    actividad (ETT-2), y resumen con la variación oficial interanual (ETT-R1..R4). No se compara con avances de hace un
    año: eran provisionales y el ministerio los revisa (p. ej. ene-jul 2025: 2.422.730 CPD en el avance, 2.551.192 revisado)."""
    url, _ = ultimo_fichero("ett")
    lector = {"mites_ett_mensual": mites_detalle.ett_mensual, "mites_ett_territorio": mites_detalle.ett_territorio,
              "mites_ett_actividad": mites_detalle.ett_actividad, "mites_ett_resumen": mites_detalle.ett_resumen}[tabla]
    cab, filas = lector(descargar_mites(url), url)
    escribir(tabla, cab, filas)
    i = cab.index("fecha") if "fecha" in cab else cab.index("periodo_hasta")
    fechas = sorted({f[i] for f in filas if f[i]})
    return [[tabla, "OK", len(filas), fechas[0], fechas[-1], url]]


def mites_atr_detalle():
    """Accidentes del último avance: por gravedad, provincia, actividad, ocupación, sexo y edad; índices por actividad y
    por CCAA y sector; variación oficial frente al mismo periodo del año anterior (ATR-R1)."""
    url, _ = ultimo_fichero("atr")
    cab, filas = mites_detalle.atr_detalle(descargar_mites(url), url)
    escribir("mites_accidentes_detalle", cab, filas)
    return [["mites_accidentes_detalle", "OK", len(filas), filas[0][0], filas[0][1], url]]


def mites_reg_detalle(meses=10):
    """Regulación de empleo: los últimos ficheros mensuales (cada uno es el dato de su mes) para tener serie.
    Los meses anteriores al último se guardan en disco (no cambian). Si una descarga falla, falla la tabla entera
    y se usa la última carga buena (mejor que escribir una serie con huecos); un mes que se descarga pero no se
    puede leer se omite y se indica en el detalle."""
    ficheros = sorted(ficheros_mites("reg").items())[-meses:]
    cab, filas, omitidos = None, [], []
    for n, ((anyo, mes), url) in enumerate(ficheros):
        datos = descargar_mites(url, fijo=n < len(ficheros) - 1)
        try:
            cab, f = mites_detalle.reg_detalle(datos, url)
            filas += f
        except Exception as e:
            omitidos.append(f"{anyo}-{mes:02d} ({type(e).__name__})")
    if not filas:
        raise ValueError(f"ningún fichero de regulación de empleo legible: {omitidos}")
    escribir("mites_regulacion_empleo", cab, filas)
    fechas = sorted({f[0] for f in filas})
    return [["mites_regulacion_empleo", "OK" if not omitidos else "AVISO", len(filas), fechas[0], fechas[-1],
             f"{len(fechas)} meses" + (f" · omitidos: {', '.join(omitidos)}" if omitidos else "")]]


def mites_cct_detalle():
    """Convenios: detalle por actividad, territorio, ámbito, jornada y cláusula de garantía (último avance y el de
    diciembre de los años anteriores desde 2023, para ver la evolución por sector y territorio)."""
    url, (anyo, mes) = ultimo_fichero("cct")
    cab, filas = mites_detalle.cct_detalle(descargar_mites(url), url)
    for a in range(2023, anyo):
        dic = ficheros_mites("cct").get((a, 12))
        if dic:
            datos = descargar_mites(dic, fijo=True)  # si la descarga falla, se usa la última carga buena completa
            try:
                filas += mites_detalle.cct_detalle(datos, dic)[1]
            except Exception as e:  # un año anterior con otra estructura no impide cargar el actual
                print(f"  CCT {a}: se omite ({type(e).__name__}: {e})")
    escribir("mites_convenios_detalle", cab, filas)
    anyos = sorted({f[0] for f in filas})
    return [["mites_convenios_detalle", "OK", len(filas), anyos[0], anyos[-1], url]]


def fundae():
    """FUNDAE: series estadísticas (pipeline/fundae.py). Cada serie conserva su última versión buena: el visor de FUNDAE
    falla a menudo y las que no se consiguen hoy se completan en las cargas siguientes."""
    cab, filas, est = fundae_lector.cargar(contexto("www.fundae.es"), NORM / "fundae")
    escribir("fundae_formacion", cab, filas)
    anyos = sorted({int(f[2]) for f in filas})
    cuenta = {k: [s for s, v in est.items() if v == k] for k in ("hoy", "anterior", "sin datos")}
    detalle = (f"{len(cuenta['hoy'])} series actualizadas hoy, {len(cuenta['anterior'])} de cargas anteriores"
               + (f", sin datos: {', '.join(cuenta['sin datos'])}" if cuenta["sin datos"] else ""))
    return [["fundae_formacion", "OK" if not cuenta["sin datos"] else "AVISO", len(filas), anyos[0], anyos[-1], detalle]]


def eurostat():
    """Eurostat (pipeline/eurostat.py): las series que fallen hoy se anotan; si falla todo, se usa la última carga buena."""
    cab, filas, fallos = eurostat_lector.cargar(descargar)
    escribir("eurostat_indicadores", cab, filas)
    anyos = sorted({f[1] for f in filas})
    detalle = "series completas" if not fallos else "; ".join(f"{k}: {v}" for k, v in fallos.items())
    return [["eurostat_indicadores", "OK" if not fallos else "AVISO", len(filas), anyos[0], anyos[-1], detalle]]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--solo", nargs="+", choices=["sepe", "ss", "mites", "fundae", "eurostat"])
    ap.add_argument("--desde", type=int, default=2019, help="primer año del SEPE por municipio (por defecto 2019)")
    a = ap.parse_args()
    QA.mkdir(parents=True, exist_ok=True)
    # (grupo, tabla que produce, función). Cada tabla se escribe solo si su lector termina bien.
    pasos = [("sepe", "sepe_paro_municipio", lambda: sepe(a.desde, "sepe_paro_municipio")),
             ("sepe", "sepe_contratos_municipio", lambda: sepe(a.desde, "sepe_contratos_municipio")),
             ("sepe", "sepe_demandantes_municipio", lambda: sepe(a.desde, "sepe_demandantes_municipio")),
             ("sepe", "sepe_demandantes_nivel", sepe_nivel_formativo),
             ("ss", "ss_afiliacion", ss), ("ss", "ss_incapacidad_temporal", ss_it), ("ss", "ss_it_detalle", ss_it_detalle),
             ("mites", "mites_cpd_actividad", mites_ett),
             ("mites", "mites_convenios", mites_cct), ("mites", "mites_accidentes", mites_atr),
             ("mites", "mites_ett_mensual", lambda: mites_ett_detalle("mites_ett_mensual")),
             ("mites", "mites_ett_territorio", lambda: mites_ett_detalle("mites_ett_territorio")),
             ("mites", "mites_ett_actividad", lambda: mites_ett_detalle("mites_ett_actividad")),
             ("mites", "mites_ett_resumen", lambda: mites_ett_detalle("mites_ett_resumen")),
             ("mites", "mites_convenios_detalle", mites_cct_detalle),
             ("mites", "mites_accidentes_detalle", mites_atr_detalle),
             ("mites", "mites_regulacion_empleo", mites_reg_detalle),
             ("fundae", "fundae_formacion", fundae),
             ("eurostat", "eurostat_indicadores", eurostat)]
    qa, ok = [], []
    for grupo, tabla, f in pasos:
        if a.solo and grupo not in a.solo:
            continue
        t0 = time.time()
        try:
            res = f()
            ok.append(tabla)
        except Exception as e:
            error = f"{type(e).__name__}: {e}"[:300]
            if (NORM / f"{tabla}.csv").exists():  # última versión buena, conservada en caché entre ejecuciones
                res = [[tabla, "ANTERIOR", "", "", "", f"{error} · se usa la última carga buena ({estado.ultima_ok(tabla)})"]]
            else:
                res = [[tabla, "ERROR", "", "", "", error]]
        for r in res:
            qa.append(r + [f"{time.time() - t0:.0f}s"])
            print(r[:5])
    estado.marcar_ok(ok)
    with open(QA / "fuentes_carga.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["tabla", "estado", "filas", "desde", "hasta", "detalle", "tiempo"]); w.writerows(qa)
    return 0


if __name__ == "__main__":
    sys.exit(main())
